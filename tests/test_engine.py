import io
import json
import pytest
from docx import Document
from resume_explorer.engine import (
    analyze, redact, shared_terms, mentioned_skills, extract_text,
    SAMPLE_JD, SAMPLE_RESUME, MAX_UPLOAD_BYTES
)


def test_demo_comparison_reproducible():
    a = analyze(SAMPLE_RESUME, SAMPLE_JD)
    b = analyze(SAMPLE_RESUME, SAMPLE_JD)
    assert a == b
    assert 'Python' in a['mentioned_in_both']
    assert 'Docker' in a['mentioned_only_in_job']
    assert 'AWS' in a['mentioned_only_in_job']
    assert 'scikit-learn' not in a['mentioned_in_both']


def test_exact_boundaries_prevent_false_positives():
    assert not mentioned_skills('I am a Reactivity engineer; I use gitter in projects.')
    assert 'React' in mentioned_skills('Built interfaces with React.')
    assert 'Git' in mentioned_skills('Used GitHub.')


def test_aliases():
    hits = mentioned_skills('My sklearn and postgres and Node.js services are tested with pytest.')
    assert {'scikit-learn','PostgreSQL','Node.js','Testing','SQL'} <= hits.keys()


def test_short_description_rejected():
    with pytest.raises(ValueError): analyze(SAMPLE_RESUME,'hi')


def test_empty_resume_rejected():
    with pytest.raises(ValueError): analyze('',SAMPLE_JD)


def test_excess_length_rejected():
    with pytest.raises(ValueError): analyze('Python ' * 8, 'A' * 50_001)


def test_redaction():
    text = 'Contact alex@example.com or +91 98765 43210 and https://example.com/x'
    result = redact(text)
    assert 'alex@example.com' not in result
    assert '98765' not in result
    assert 'example.com' not in result
    assert '[email redacted]' in result


def test_evidence_is_best_effort_redacted():
    a = analyze('alex@example.com\n' + ('Python FastAPI SQL experience. ' * 4), SAMPLE_JD)
    assert 'alex@example.com' not in json.dumps(a)


def test_shared_terms_present():
    assert 'python' in shared_terms('Built python python systems with data', 'python data systems skills needed')


def test_shared_terms_empty():
    assert shared_terms('','job') == []
    assert shared_terms('and the','at of') == []


def test_txt_extract():
    text = 'Python SQL analysis and reporting for fictional demonstration.'
    assert extract_text(text.encode(),'demo.txt') == text


def test_oversize_rejected():
    with pytest.raises(ValueError): extract_text(b'x'*(MAX_UPLOAD_BYTES+1),'big.txt')


def test_bad_pdf_header_rejected():
    with pytest.raises(ValueError): extract_text(b'not a pdf at all','fake.pdf')


def test_bad_extension_rejected():
    with pytest.raises(ValueError): extract_text(b'some data here','file.exe')


def test_docx_extract():
    document = Document()
    document.add_paragraph('This fictional resume describes Python and FastAPI analytics.')
    table = document.add_table(rows=1, cols=2)
    table.cell(0,0).text='SQL'
    table.cell(0,1).text='Pandas'
    buf=io.BytesIO()
    document.save(buf)
    result=extract_text(buf.getvalue(),'demo.docx')
    assert 'Python' in result and 'SQL | Pandas' in result


def test_scanned_pdf_warning():
    with pytest.raises(ValueError): extract_text(b'%PDF-not-pdf','fake.pdf')


def test_no_hiring_recommendations():
    result=analyze(SAMPLE_RESUME,SAMPLE_JD)
    assert all(x not in result for x in ('ats_score','hiring_confidence','recommendation','shortlist'))


def test_pdf_extract_generated_with_real_parser():
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    page = canvas.Canvas(buf, pagesize=(500, 600))
    page.drawString(40, 530, 'Fictional resume with Python, SQL, FastAPI and Pandas.')
    page.save()
    assert 'Python' in extract_text(buf.getvalue(), 'resume.pdf')


def test_api_without_personal_scoring():
    from fastapi.testclient import TestClient
    from api.main import app
    client = TestClient(app)
    assert client.get('/health').json() == {'ok': True, 'demo': True, 'external_ai': False}
    response = client.post('/api/explore', json={'resume_text': SAMPLE_RESUME, 'job_description': SAMPLE_JD})
    assert response.status_code == 200
    assert 'mentioned_in_both' in response.json()
    assert 'ats_score' not in response.json()


def test_api_invalid_input_422():
    from fastapi.testclient import TestClient
    from api.main import app
    client = TestClient(app)
    assert client.post('/api/explore', json={'resume_text': 'short', 'job_description': 'x'}).status_code == 422
